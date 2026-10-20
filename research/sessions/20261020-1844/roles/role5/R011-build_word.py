from pathlib import Path
import re, json, hashlib, copy, sys, argparse
from urllib.parse import unquote
from lxml import etree
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ap=argparse.ArgumentParser(description='Frozen Markdown to editable DOCX; source hash is mandatory.')
ap.add_argument('--source',required=True,type=Path)
ap.add_argument('--sha256',required=True)
ap.add_argument('--output',required=True,type=Path)
ap.add_argument('--report',type=Path)
ap.add_argument('--check-source',action='store_true',help='Parse and validate only; never create DOCX.')
args=ap.parse_args()
SOURCE=args.source.resolve(); OUTPUT=args.output.resolve(); EXPECTED=args.sha256.lower()
if not re.fullmatch('[0-9a-f]{64}',EXPECTED):ap.error('invalid SHA256')
if hashlib.sha256(SOURCE.read_bytes()).hexdigest()!=EXPECTED:raise ValueError('source SHA mismatch')
TEXT=SOURCE.read_text('utf-8')
if OUTPUT==SOURCE:raise ValueError('output must differ from source')
if OUTPUT.suffix.lower()!='.docx':raise ValueError('output must be .docx')
WORK=OUTPUT.parent
def el(name, **attrs):
    node = OxmlElement(name)
    for k,v in attrs.items(): node.set(qn(k), str(v))
    return node

def mathrun(text, plain=False):
    r = el('m:r')
    rp = el('m:rPr'); rp.append(el('m:sty', **{'m:val':'p' if plain else 'i'})); r.append(rp)
    wrp = el('w:rPr'); wrp.append(el('w:rFonts', **{'w:ascii':'Cambria Math','w:hAnsi':'Cambria Math','w:eastAsia':'SimSun'})); wrp.append(el('w:sz', **{'w:val':'21'})); r.append(wrp)
    t = el('m:t'); t.text = text; r.append(t)
    return r

def container(name, nodes):
    node=el(name)
    for x in nodes: node.append(x)
    return node

SYMBOLS = dict(kappa='κ',delta='δ',alpha='α',chi='χ',rho='ρ',lambda_='λ',gamma='γ',varepsilon='ε',zeta='ζ',phi='φ',omega='ω',pi='π',Delta='Δ',cdot='·',le='≤',ge='≥',in_='∈',sum='∑')
SYMBOLS['lambda']=SYMBOLS.pop('lambda_'); SYMBOLS['in']=SYMBOLS.pop('in_')

class Latex:
    """Strict LaTeX subset to editable OMML; unsupported syntax fails before writing."""
    def __init__(self, s): self.s=s; self.i=0
    def skip(self):
        while self.i<len(self.s) and self.s[self.i].isspace(): self.i+=1
    def argument(self):
        self.skip()
        if self.i<len(self.s) and self.s[self.i]=='{':
            self.i+=1; result=self.sequence('}'); return result
        return self.atom()
    def sequence(self, stop=None):
        out=[]
        while self.i<len(self.s):
            self.skip()
            if self.i>=len(self.s): break
            if stop and self.s[self.i]==stop: self.i+=1; return out
            if self.s[self.i]=='}': raise ValueError('unmatched brace')
            base=self.atom(); self.skip(); sub=sup=None
            while self.i<len(self.s) and self.s[self.i] in '_^':
                kind=self.s[self.i]; self.i+=1; arg=self.argument()
                if kind=='_': sub=arg
                else: sup=arg
                self.skip()
            if sub is not None or sup is not None:
                text=''.join(n.text or '' for b in base for n in b.iter(qn('m:t')))
                if sub is not None and text in ['∑','min','max']:
                    x=el('m:limLow'); x.append(container('m:e',base)); x.append(container('m:lim',sub)); base=[x]; sub=None
                if sub is not None and sup is not None:
                    x=el('m:sSubSup'); x.append(container('m:e',base)); x.append(container('m:sub',sub)); x.append(container('m:sup',sup)); base=[x]
                elif sub is not None:
                    isbrace=any(n.tag==qn('m:groupChr') for b in base for n in b.iter())
                    x=el('m:limLow' if isbrace else 'm:sSub'); x.append(container('m:e',base)); x.append(container('m:lim' if isbrace else 'm:sub',sub)); base=[x]
                elif sup is not None:
                    x=el('m:sSup'); x.append(container('m:e',base)); x.append(container('m:sup',sup)); base=[x]
            out.extend(base)
        if stop: raise ValueError('unclosed '+stop)
        return out
    def atom(self):
        self.skip()
        if self.i>=len(self.s): return []
        c=self.s[self.i]; self.i+=1
        if c=='{': return self.sequence('}')
        if c!='\\': return [mathrun(c, c.isdigit() or c in '+-=(),[]|./:')]
        if self.i>=len(self.s): raise ValueError('trailing slash')
        if not self.s[self.i].isalpha():
            c=self.s[self.i]; self.i+=1
            if c in '!': return []
            if c in ',; ': return [mathrun('\u2009',True)]
            if c in '{}': return [mathrun(c,True)]
            raise ValueError('unknown escaped '+c)
        start=self.i
        while self.i<len(self.s) and self.s[self.i].isalpha(): self.i+=1
        cmd=self.s[start:self.i]
        if cmd in ['Phi','ell','mu','sigma','theta','Gamma','Theta']:
            return [mathrun({'Phi':'Φ','ell':'ℓ','mu':'μ','sigma':'σ','theta':'θ','Gamma':'Γ','Theta':'Θ'}[cmd])]
        if cmd in ['sin','cos','tan','diag','sgn']: return [mathrun(cmd,True)]
        if cmd in ['left','right']:
            self.skip()
            if self.i < len(self.s) and self.s[self.i]=='.': self.i+=1; return []
            return self.atom()
        if cmd=='begin':
            self.skip()
            if self.s[self.i]!='{': raise ValueError('missing environment')
            end=self.s.index('}',self.i); env=self.s[self.i+1:end]; self.i=end+1
            if env not in ['cases','matrix','pmatrix','bmatrix']: raise ValueError('unsupported environment '+env)
            marker='\\end{'+env+'}'; end=self.s.index(marker,self.i)
            body=self.s[self.i:end]; self.i=end+len(marker)
            rows=[x for x in re.split(r'\\\\',body) if x.strip()]
            matrix=el('m:m')
            for row in rows:
                mr=el('m:mr')
                for cell in row.split('&'): mr.append(container('m:e',Latex(cell).sequence()))
                matrix.append(mr)
            if env=='matrix': return [matrix]
            d=el('m:d'); pr=el('m:dPr')
            beg,fin={'cases':('{',''),'pmatrix':('(',')'),'bmatrix':('[',']')}[env]
            pr.append(el('m:begChr',**{'m:val':beg}));pr.append(el('m:endChr',**{'m:val':fin}))
            d.append(pr);d.append(container('m:e',[matrix]));return [d]
        if cmd in SYMBOLS: return [mathrun(SYMBOLS[cmd],cmd in ['sum','cdot','le','ge','in'])]
        if cmd in ['min','max']: return [mathrun(cmd,True)]
        if cmd in ['quad','qquad']: return [mathrun('\u2003',True)]
        if cmd=='frac':
            a=self.argument(); b=self.argument(); x=el('m:f'); x.append(container('m:num',a)); x.append(container('m:den',b)); return [x]
        if cmd=='dot':
            x=el('m:acc'); p=el('m:accPr'); p.append(el('m:chr',**{'m:val':'\u0307'})); x.append(p); x.append(container('m:e',self.argument())); return [x]
        if cmd in ['operatorname','mathrm','text','boldsymbol','mathcal']:
            nodes=self.argument()
            for node in nodes:
                for r in node.iter(qn('m:r')):
                    rp=r.find(qn('m:rPr'))
                    if rp is None: rp=el('m:rPr');r.insert(0,rp)
                    for old in list(rp): rp.remove(old)
                    if cmd=='mathcal': rp.append(el('m:scr',**{'m:val':'script'})); rp.append(el('m:sty',**{'m:val':'p'}))
                    else: rp.append(el('m:sty',**{'m:val':'bi' if cmd=='boldsymbol' else 'p'}))
            return nodes
        if cmd=='underbrace':
            x=el('m:groupChr'); p=el('m:groupChrPr'); p.append(el('m:chr',**{'m:val':'\u23df'})); p.append(el('m:pos',**{'m:val':'bot'})); p.append(el('m:vertJc',**{'m:val':'bot'})); x.append(p); x.append(container('m:e',self.argument())); return [x]
        raise ValueError('unsupported command '+cmd)

def parse_math(s):
    parser=Latex(s); nodes=parser.sequence(); assert parser.i==len(s); return nodes


# Known literal arrays are presentation-transcribed without changing their values.
TRANSFORMS={
'M=diag(1,1.2,.8)':r'$M=\operatorname{diag}(1,1.2,.8)$',
'K=[[2.4,−1,0],[−1,1.7,−.7],[0,−.7,.7]]':r'$K=\begin{bmatrix}2.4&-1&0\\-1&1.7&-.7\\0&-.7&.7\end{bmatrix}$',
'B=(1,0,0)ᵀ':r'$B=(1,0,0)^T$',
'ΦᵀMΦ=I':r'$\Phi^TM\Phi=I$',
'u=Φq':r'$u=\Phi q$',
'q=ΦᵀMu':r'$q=\Phi^TMu$',
'F_IS=T_SK u':r'$F_{IS}=T_SKu$',
'ℓsᵀ=e_sᵀT_SK':r'$\ell_s^T=e_s^TT_SK$',
'T_S(C udot+Ku)':r'$T_S(C\dot u+Ku)$',
'g(F)=sgn(F)(|F|/c)^(1/α)':r'$g(F)=\operatorname{sgn}(F)(|F|/c)^{1/\alpha}$',
'V2,ref(t)=Σa_rV2,r,0(t)':r'$V_{2,\mathrm{ref}}(t)=\sum_r a_r V_{2,r,0}(t)$',

}
CONTENT=TEXT
for a,b in TRANSFORMS.items():CONTENT=CONTENT.replace(a,b)
excluded=[]
kept=[]
for line in CONTENT.splitlines():
    if line.startswith('完整计算条件、原文定位和资料范围见'):
        excluded.append(line);continue
    kept.append(line)
CONTENT='\n'.join(kept)
raw_equations=re.findall(r'\$\$(.*?)\$\$',CONTENT,re.S)
parsed_displays=[]
for raw in raw_equations:
    clean=re.sub(r'\\tag\{\d+\}','',raw).strip()
    nodes=parse_math(clean)
    parsed_displays.append({'source':clean,'OMML_nodes':len(nodes)})
without_display=re.sub(r'\$\$.*?\$\$','',CONTENT,flags=re.S)
inline_equations=re.findall(r'\$([^$]+)\$',without_display)
for raw in inline_equations:parse_math(raw)
images=[]
for target in re.findall(r'!\[[^]]*\]\(([^)]+)\)',CONTENT):
    image=(SOURCE.parent/unquote(target)).resolve()
    if not image.is_file():raise FileNotFoundError(image)
    images.append({'path':str(image),'sha256':hashlib.sha256(image.read_bytes()).hexdigest()})
if len(images)>3:raise ValueError('contract permits at most 3 main figures')
table_groups=re.findall(r'(?m)^\|.*(?:\n\|.*)*',CONTENT)
for group in table_groups:
    widths=[len(l.strip('|').split('|')) for l in group.splitlines()]
    if len(set(widths))!=1:raise ValueError('ragged source table')
# No commands may escape a math block into Word body text.
plain=re.sub(r'\$[^$]*\$','',without_display)
if re.search(r'\\[A-Za-z]+',plain):raise ValueError('LaTeX outside math block')
if '[[' in plain:raise ValueError('unconverted program array')
check={'mode':'check-source' if args.check_source else 'build','source':str(SOURCE),'source_sha256':EXPECTED,
       'display_equations':parsed_displays,'inline_math_groups':len(inline_equations),'tables':len(table_groups),
       'images':images,'references':len(re.findall(r'^\[\d+\]',CONTENT,re.M)),
       'excluded_process_tail':excluded,'native_math_only':True,'source_unchanged':True}
if args.check_source:
    if args.report:
        if args.report.exists():raise FileExistsError(args.report)
        args.report.write_text(json.dumps(check,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(check,ensure_ascii=False,indent=2));sys.exit(0)
if OUTPUT.exists():raise FileExistsError(OUTPUT)
if (args.report or OUTPUT.with_suffix('.build-report.json')).exists():raise FileExistsError('report exists')
if not OUTPUT.parent.is_dir():raise FileNotFoundError(OUTPUT.parent)
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(1.9); sec.left_margin=Cm(2.2); sec.right_margin=Cm(2.2)
sec.header_distance=Cm(.8); sec.footer_distance=Cm(.8)

def fontset(obj, size, east='SimSun', bold=False):
    obj.font.name='Times New Roman'; obj.font.size=Pt(size); obj.font.bold=bold; obj.font.italic=False; obj.font.color.rgb=RGBColor(0,0,0)
    prop=obj._element.get_or_add_rPr()
    f=prop.find(qn('w:rFonts'))
    if f is None: f=el('w:rFonts');prop.insert(0,f)
    for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',east)]: f.set(qn('w:'+k),v)

for name,size,east,bold in [('Normal',11,'SimSun',False),('Title',18,'SimHei',True),('Subtitle',11.5,'SimSun',True),('Heading 1',13.5,'SimHei',True),('Heading 2',11.5,'SimHei',True),('Caption',10,'SimSun',True)]:
    style=doc.styles[name];fontset(style,size,east,bold)
    pf=style.paragraph_format; pf.space_before=Pt(0);pf.space_after=Pt(4);pf.line_spacing=1.20
    if name.startswith('Heading'): pf.space_before=Pt(10);pf.space_after=Pt(6);pf.keep_with_next=True;pf.line_spacing=1.1
    if name in ['Title','Subtitle']: pf.space_after=Pt(8);pf.line_spacing=1.1
normal=doc.styles['Normal']; normal.paragraph_format.widow_control=True
for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
mathpr=el('m:mathPr'); mathpr.append(el('m:mathFont',**{'m:val':'Cambria Math'})); doc.settings.element.append(mathpr)
doc.core_properties.title=CONTENT.splitlines()[0].removeprefix('# ')
doc.core_properties.author='';doc.core_properties.subject='楼层剪力分析论文审阅稿'
doc.core_properties.comments=''

head=sec.header.paragraphs[0];head.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=head.add_run(CONTENT.splitlines()[0].removeprefix('# '));fontset(r,8.5)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text,field in [('第 ',None),(None,'PAGE'),(' 页',None)]:
    r=footer.add_run();fontset(r,9)
    if text:r.text=text
    else:
        fld=el('w:fldSimple',**{'w:instr':field}); rr=el('w:r');tt=el('w:t');tt.text='1';rr.append(tt);fld.append(rr);r._r.addnext(fld)

def hyperlink(p,label,target,size=None):
    resolved=target if re.match(r'^[a-z]+:',target,re.I) else (SOURCE.parent/unquote(target)).resolve().as_uri()
    h=el('w:hyperlink',**{'r:id':doc.part.relate_to(resolved,RT.HYPERLINK,is_external=True)})
    r=el('w:r');rp=el('w:rPr');rp.append(el('w:color',**{'w:val':'1F4E79'}));rp.append(el('w:rFonts',**{'w:ascii':'Times New Roman','w:hAnsi':'Times New Roman','w:eastAsia':'SimSun'}))
    if size:rp.append(el('w:sz',**{'w:val':str(int(size*2))}))
    r.append(rp); t=el('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

INLINE=re.compile(r'(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\)|\$[^$]+\$)')
def inline(p,s,size=None):
    pos=0
    for m in INLINE.finditer(s):
        if m.start()>pos:
            r=p.add_run(s[pos:m.start()]);
            if size:fontset(r,size)
        token=m.group()
        if token.startswith('**'):
            r=p.add_run(token[2:-2]);r.bold=True
            if size:fontset(r,size,bold=True)
        elif token.startswith('$'):
            p._p.append(container('m:oMath',parse_math(token[1:-1])))
        else:
            mm=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token);hyperlink(p,*mm.groups(),size)
        pos=m.end()
    if pos<len(s):
        r=p.add_run(s[pos:]);
        if size:fontset(r,size)

def para(s,english=False,note=False):
    p=doc.add_paragraph();pf=p.paragraph_format
    pf.first_line_indent=Cm(0) if english or note else Pt(22)
    if english:pf.line_spacing=1.15
    inline(p,s,9.2 if note else (10.5 if english else None))
    return p

equation_log=[]
def equation(s):
    match=re.search(r'\\tag\{(\d+)\}',s); tag=match.group(1) if match else None; s=re.sub(r'\\tag\{\d+\}','',s).strip()
    parts=[x.strip() for x in re.split(r'\\(?:qquad|quad)\b',s) if x.strip()]
    # Keep the two H2 factors readable without altering factors, phase or coefficients.
    if s.startswith('a_g(t)=') and r']\,[' in s:
        left,right=s.split(r']\,[',1)
        parts=[left+']','·['+right]
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_together=True
    tabs=p.paragraph_format.tab_stops;tabs.add_tab_stop(Cm(7.65),WD_TAB_ALIGNMENT.CENTER);tabs.add_tab_stop(Cm(16.5),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t'); om=el('m:oMath')
    if len(parts)>1:
        arr=el('m:eqArr');pr=el('m:eqArrPr');pr.append(el('m:baseJc',**{'m:val':'center'}));arr.append(pr)
        for part in parts:arr.append(container('m:e',parse_math(part)))
        om.append(arr)
    else:
        for x in parse_math(s):om.append(x)
    p._p.append(om)
    if tag:
        r=p.add_run('\t('+tag+')');fontset(r,10.5)
    equation_log.append({'number':tag,'source_latex':s,'display_lines':len(parts),'representation':'native OMML'})

table_log=[]
def add_table(lines):
    rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    rows.pop(1);n=len(rows[0]);assert all(len(x)==n for x in rows)
    idx=len(table_log)+1
    presets={1:[1.1,1.1,1.3,1.4,1.4],2:[1.2,1,1,1,1,1,1],3:[1,1,1,1,1,1],
             4:[.6,1.5,1.5,1.5,1.7],5:[2,1,1.5,1.5,1.5,1.5],
             6:[1,1.5,1.5,1.5],7:[1,1.6,1.5,1.5,1.5,1.5],8:[2.8,1.5,1.1,1.6]}
    weights=presets.get(idx,[1]*n)
    if len(weights)!=n:weights=[max(1,min(3,len(rows[0][j])/5)) for j in range(n)]
    widths=[16.6*w/sum(weights) for w in weights]
    table=doc.add_table(rows=1,cols=n);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    props=table._tbl.tblPr
    borders=el('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:borders.append(el('w:'+edge,**{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
    props.append(borders)
    margins=el('w:tblCellMar')
    for edge in ['top','bottom','left','right']:margins.append(el('w:'+edge,**{'w:w':'75','w:type':'dxa'}))
    props.append(margins)
    for j,w in enumerate(widths):table.columns[j].width=Cm(w)
    for i,row in enumerate(rows):
        cells=table.rows[0].cells if i==0 else table.add_row().cells
        trpr=table.rows[i]._tr.get_or_add_trPr();trpr.append(el('w:cantSplit'))
        if i==0:trpr.append(el('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:cell._tc.get_or_add_tcPr().append(el('w:shd',**{'w:fill':'EDEFF2'}))
            cp=cell.paragraphs[0];cp.alignment=WD_ALIGN_PARAGRAPH.LEFT if (idx in [5,6] and j==0) or (idx==6 and j>0) else WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.line_spacing=1.08;cp.paragraph_format.space_before=Pt(2);cp.paragraph_format.space_after=Pt(2);cp.paragraph_format.keep_with_next=(i==0)
            inline(cp,value,8.8 if idx in [1,3] else 9.2)
            if i==0:
                for r in cp.runs:r.bold=True
    table_log.append({'number':idx,'rows':len(rows),'columns':n,'source_cells':rows})
    after=doc.add_paragraph();after.paragraph_format.space_after=Pt(0);after.paragraph_format.space_before=Pt(0);after.paragraph_format.line_spacing=Pt(2);after.add_run().font.size=Pt(2)

lines=CONTENT.splitlines(); i=0;english=False;paragraph_log=[];image_log=[]
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line.startswith('$$'):
        block=line
        while not (block.endswith('$$') and len(block)>4):
            if i>=len(lines):raise ValueError('unclosed display')
            block+='\n'+lines[i];i+=1
        equation(block[2:-2]);continue
    if line.startswith('|'):
        group=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i].strip());i+=1
        add_table(group);continue
    if line.startswith('!['):
        m=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line); label,target=m.groups();img=(SOURCE.parent/target).resolve();assert img.is_file()
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(img),width=Cm(14.0))
        
        image_log.append({'source':str(img),'caption':label});continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;continue
    if line.startswith('## '):
        title=line[3:];english=title=='Abstract';doc.add_heading(title,level=1);continue
    if line.startswith('### '):doc.add_heading(line[4:],level=2);continue
    if line.startswith('**Modal-superposition') or line.startswith('**Placement of nonlinear'):
        p=doc.add_paragraph(line[2:-2],style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;continue
    if line.startswith('作者') or line.startswith('Authors:'):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;inline(p,line,10);continue
    if line.startswith('> '):para(line[2:],note=True);continue
    if re.match(r'^图\d+\s',line):
        pr=doc.add_paragraph(line,style='Caption');pr.alignment=WD_ALIGN_PARAGRAPH.CENTER;pr.paragraph_format.keep_with_next=False
        continue
    if line.startswith('**表'):
        p=doc.add_paragraph(line[2:-2],style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;continue
    if re.match(r'^\[\d+\]',line):
        p=para(line,english=True);p.paragraph_format.left_indent=Pt(17);p.paragraph_format.first_line_indent=Pt(-17);p.paragraph_format.space_after=Pt(5)
        for r in p.runs:fontset(r,9.5)
    elif line.startswith('**关键词') or line.startswith('**Keywords'):
        p=para(line,english=True);p.paragraph_format.space_after=Pt(8)
    else:
        p=para(line,english=english)
        if re.match(r'^（[1-4]）',line):p.paragraph_format.keep_together=True
        nextline=next((x.strip() for x in lines[i:] if x.strip()),'')
        if nextline.startswith('$$'):p.paragraph_format.keep_with_next=True
    paragraph_log.append(line)


assert len(equation_log)==len(raw_equations)
assert len(table_log)==len(table_groups) and len(image_log)==len(images)
doc.save(OUTPUT)
# Inspect the serialized package for visible LaTeX leaks and native equation structures.
import zipfile
with zipfile.ZipFile(OUTPUT) as z:
    root=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    visible=''.join(root.xpath('//w:t/text() | //m:t/text()',namespaces=ns))
    if re.search(r'\\[A-Za-z]+|\$\$',visible):raise ValueError('visible LaTeX leak')
    check['serialized_math_groups']=len(root.xpath('//m:oMath',namespaces=ns))
    check['serialized_matrices']=len(root.xpath('//m:m',namespaces=ns))
check.update({'output':str(OUTPUT),'output_sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
              'equations':equation_log,'table_details':table_log,'image_details':image_log,
              'source_unchanged':hashlib.sha256(SOURCE.read_bytes()).hexdigest()==EXPECTED})
report=args.report or OUTPUT.with_suffix('.build-report.json')
if report.exists():raise FileExistsError(report)
report.write_text(json.dumps(check,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'output':str(OUTPUT),'sha256':check['output_sha256'],'math_groups':check['serialized_math_groups'],
                  'tables':len(table_log),'images':len(image_log)},ensure_ascii=False))
